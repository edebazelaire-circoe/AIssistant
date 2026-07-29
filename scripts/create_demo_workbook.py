from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


STATUS = {
    "À faire": "F4CCCC",
    "En cours": "FFF2CC",
    "Terminé": "D9EAD3",
}


def create(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Roadmap 2026"
    headers = ["ID", "Action", "Responsable", "Mois", "Statut", "Note"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="D9E2F3")
    rows = [
        ["A-001", "Sécuriser le portail client", "Alice", "Juillet", "En cours", "Priorité haute"],
        ["A-002", "Mettre à jour la documentation API", "Bob", "Août", "À faire", ""],
        ["A-003", "Valider le prototype Jarvis", "Étienne", "Septembre", "À faire", ""],
    ]
    for row in rows:
        sheet.append(row)
    for row_number in range(2, sheet.max_row + 1):
        status_cell = sheet.cell(row_number, 5)
        status_cell.fill = PatternFill(fill_type="solid", fgColor=STATUS[status_cell.value])
    sheet["H1"] = "Cellule hors table à préserver"
    sheet["H2"] = "=COUNTA(A2:A20)"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:F{sheet.max_row}"
    widths = {"A": 12, "B": 38, "C": 18, "D": 14, "E": 16, "F": 28}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    workbook.save(path)
    return path


if __name__ == "__main__":
    target = Path(__file__).resolve().parents[1] / "workspace" / "roadmap_demo.xlsx"
    print(create(target))
