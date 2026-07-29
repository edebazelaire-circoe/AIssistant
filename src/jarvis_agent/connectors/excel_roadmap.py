from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
import unicodedata
from dataclasses import asdict, dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from jarvis_agent.domain.models import (
    Artifact,
    ArtifactDiff,
    Diagnostic,
    ErrorCode,
    ResultStatus,
    RoadmapItem,
    RoadmapStatus,
    ToolResult,
)


STATUS_ALIASES: dict[str, RoadmapStatus] = {
    "not done": RoadmapStatus.NOT_DONE,
    "not_done": RoadmapStatus.NOT_DONE,
    "todo": RoadmapStatus.NOT_DONE,
    "to do": RoadmapStatus.NOT_DONE,
    "a faire": RoadmapStatus.NOT_DONE,
    "à faire": RoadmapStatus.NOT_DONE,
    "non fait": RoadmapStatus.NOT_DONE,
    "in progress": RoadmapStatus.IN_PROGRESS,
    "in_progress": RoadmapStatus.IN_PROGRESS,
    "doing": RoadmapStatus.IN_PROGRESS,
    "en cours": RoadmapStatus.IN_PROGRESS,
    "done": RoadmapStatus.DONE,
    "complete": RoadmapStatus.DONE,
    "completed": RoadmapStatus.DONE,
    "termine": RoadmapStatus.DONE,
    "terminé": RoadmapStatus.DONE,
    "fait": RoadmapStatus.DONE,
}

STATUS_LABELS: dict[RoadmapStatus, str] = {
    RoadmapStatus.NOT_DONE: "À faire",
    RoadmapStatus.IN_PROGRESS: "En cours",
    RoadmapStatus.DONE: "Terminé",
}

STATUS_COLORS: dict[RoadmapStatus, str] = {
    RoadmapStatus.NOT_DONE: "F4CCCC",
    RoadmapStatus.IN_PROGRESS: "FFF2CC",
    RoadmapStatus.DONE: "D9EAD3",
}

ACTION_HEADERS = {"action", "actions", "task", "tache", "tâche", "description", "work item"}
STATUS_HEADERS = {"status", "statut", "etat", "état", "avancement"}
OWNER_HEADERS = {"owner", "responsable", "pilote"}
MONTH_HEADERS = {"month", "mois", "periode", "période"}
ID_HEADERS = {"id", "key", "cle", "clé", "action id"}


@dataclass(slots=True, frozen=True)
class RoadmapSchema:
    sheet_name: str
    header_row: int
    action_col: int
    status_col: int
    owner_col: int | None = None
    month_col: int | None = None
    id_col: int | None = None


@dataclass(slots=True, frozen=True)
class RowMatch:
    item: RoadmapItem
    score: float


class ExcelRoadmapConnector:
    id = "excel"

    def __init__(self, allowed_root: Path) -> None:
        self.allowed_root = allowed_root.resolve()
        self.allowed_root.mkdir(parents=True, exist_ok=True)

    def health(self) -> dict[str, Any]:
        writable = os.access(self.allowed_root, os.W_OK)
        return {
            "connector_id": self.id,
            "status": "healthy" if writable else "degraded",
            "message": f"Allowed root: {self.allowed_root}",
            "permissions": {"allowed_roots": [str(self.allowed_root)], "writable": writable},
        }

    def resolve_allowed(self, path: str | Path) -> Path:
        candidate = Path(path)
        if not candidate.is_absolute():
            candidate = self.allowed_root / candidate
        candidate = candidate.expanduser().resolve()
        try:
            candidate.relative_to(self.allowed_root)
        except ValueError as exc:
            raise PermissionError(f"Path is outside allowed root: {candidate}") from exc
        return candidate

    def discover(self) -> list[Path]:
        return sorted(
            path for path in self.allowed_root.rglob("*.xlsx") if not path.name.startswith("~$")
        )

    @staticmethod
    def fingerprint(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    def detect_schema(self, path: Path, sheet_hint: str | None = None) -> RoadmapSchema:
        workbook = load_workbook(path, read_only=True, data_only=False)
        candidates: list[RoadmapSchema] = []
        for sheet in workbook.worksheets:
            if sheet_hint and sheet_hint.lower() not in sheet.title.lower():
                continue
            for row_number in range(1, min(sheet.max_row, 12) + 1):
                headers: dict[str, int] = {}
                for cell in sheet[row_number]:
                    key = _normalize(str(cell.value or ""))
                    if key:
                        headers[key] = cell.column
                action_col = _find_header(headers, ACTION_HEADERS)
                status_col = _find_header(headers, STATUS_HEADERS)
                if action_col and status_col:
                    candidates.append(
                        RoadmapSchema(
                            sheet_name=sheet.title,
                            header_row=row_number,
                            action_col=action_col,
                            status_col=status_col,
                            owner_col=_find_header(headers, OWNER_HEADERS),
                            month_col=_find_header(headers, MONTH_HEADERS),
                            id_col=_find_header(headers, ID_HEADERS),
                        )
                    )
        workbook.close()
        if len(candidates) != 1:
            raise ValueError(
                f"Expected one roadmap schema, found {len(candidates)}: "
                f"{[item.sheet_name for item in candidates]}"
            )
        return candidates[0]

    def read_items(self, path: Path, schema: RoadmapSchema) -> list[RoadmapItem]:
        workbook = load_workbook(path, read_only=False, data_only=False)
        sheet = workbook[schema.sheet_name]
        items: list[RoadmapItem] = []
        for row_number in range(schema.header_row + 1, sheet.max_row + 1):
            action = str(sheet.cell(row_number, schema.action_col).value or "").strip()
            if not action:
                continue
            raw_status = str(sheet.cell(row_number, schema.status_col).value or "").strip()
            status = parse_status(raw_status)
            if status is None:
                # Unrecognized values are intentionally not mutated.
                continue
            owner = (
                str(sheet.cell(row_number, schema.owner_col).value or "").strip()
                if schema.owner_col
                else None
            )
            month = (
                str(sheet.cell(row_number, schema.month_col).value or "").strip()
                if schema.month_col
                else None
            )
            row_id = (
                str(sheet.cell(row_number, schema.id_col).value or "").strip()
                if schema.id_col
                else f"{schema.sheet_name}:{row_number}"
            )
            fill = sheet.cell(row_number, schema.status_col).fill
            color = fill.fgColor.rgb[-6:] if fill and fill.fgColor and fill.fgColor.rgb else ""
            items.append(
                RoadmapItem(
                    row_id=row_id,
                    action_text=action,
                    owner=owner or None,
                    month=month or None,
                    status=status,
                    status_cell=f"{get_column_letter(schema.status_col)}{row_number}",
                    display_color=color,
                    source_sheet=schema.sheet_name,
                    source_row_number=row_number,
                )
            )
        workbook.close()
        return items

    def match_row(
        self,
        items: list[RoadmapItem],
        action_text: str,
        owner: str | None = None,
        month: str | None = None,
        threshold: float = 0.62,
        margin: float = 0.08,
    ) -> RowMatch:
        scored: list[RowMatch] = []
        for item in items:
            score = action_similarity(action_text, item.action_text)
            if owner and item.owner:
                score += 0.08 if _normalize(owner) == _normalize(item.owner) else -0.05
            if month and item.month:
                score += 0.06 if _normalize(month) == _normalize(item.month) else -0.03
            scored.append(RowMatch(item=item, score=max(0.0, min(1.0, score))))
        scored.sort(key=lambda match: match.score, reverse=True)
        if not scored or scored[0].score < threshold:
            raise LookupError("No roadmap row reached the match threshold")
        if len(scored) > 1 and scored[0].score - scored[1].score < margin:
            raise RuntimeError(
                f"Ambiguous row match: {scored[0].item.action_text!r} ({scored[0].score:.2f}) "
                f"and {scored[1].item.action_text!r} ({scored[1].score:.2f})"
            )
        return scored[0]

    def propose(
        self,
        path: Path,
        action_text: str,
        desired_status: RoadmapStatus,
        owner: str | None = None,
        month: str | None = None,
        sheet_hint: str | None = None,
    ) -> dict[str, Any]:
        schema = self.detect_schema(path, sheet_hint)
        items = self.read_items(path, schema)
        matched = self.match_row(items, action_text, owner, month)
        before = matched.item
        return {
            "workbook_path": str(path),
            "workbook_fingerprint": self.fingerprint(path),
            "schema": asdict(schema),
            "match": {"score": matched.score, "item": before.model_dump(mode="json")},
            "changes": [
                {
                    "sheet": before.source_sheet,
                    "cell": before.status_cell,
                    "before": STATUS_LABELS[before.status],
                    "after": STATUS_LABELS[desired_status],
                    "before_status": before.status.value,
                    "after_status": desired_status.value,
                    "before_color": before.display_color,
                    "after_color": STATUS_COLORS[desired_status],
                }
            ],
            "already_applied": before.status == desired_status,
        }

    def apply(self, proposal: dict[str, Any]) -> ToolResult:
        try:
            path = self.resolve_allowed(proposal["workbook_path"])
        except PermissionError as exc:
            return _failure(ErrorCode.PATH_NOT_ALLOWED, str(exc), "excel.apply")
        if not path.exists():
            return _failure(ErrorCode.ROADMAP_ROW_NOT_FOUND, f"Workbook not found: {path}", "excel.apply")
        current_fingerprint = self.fingerprint(path)
        if current_fingerprint != proposal["workbook_fingerprint"]:
            return _failure(
                ErrorCode.FILE_CHANGED,
                "Workbook changed after preview; refusing to overwrite",
                "excel.apply",
                retryable=True,
            )
        change = proposal["changes"][0]
        desired_status = RoadmapStatus(change["after_status"])
        before_version = current_fingerprint
        if proposal.get("already_applied"):
            artifact_id = f"artifact_{uuid4().hex}"
            diff_id = f"diff_{uuid4().hex}"
            artifact = Artifact(
                id=artifact_id,
                artifact_type="excel_workbook",
                display_name=path.name,
                uri_or_local_ref=str(path),
                mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                preview_ref=f"preview:{artifact_id}",
                version=current_fingerprint,
            )
            no_op_change = {**change, "no_op": True}
            diff = ArtifactDiff(
                id=diff_id,
                artifact_id=artifact_id,
                before_version=current_fingerprint,
                after_version=current_fingerprint,
                changes=(no_op_change,),
                visual_highlights=(
                    {
                        "sheet": change["sheet"],
                        "cell": change["cell"],
                        "color": change["after_color"],
                    },
                ),
            )
            return ToolResult(
                status=ResultStatus.SUCCESS,
                summary="Roadmap already has the requested status; no file write was needed",
                data={
                    "artifact": artifact.model_dump(mode="json"),
                    "diff": diff.model_dump(mode="json"),
                    "preview": self.preview(path, change["sheet"], [change["cell"]]),
                    "workbook_path": str(path),
                    "no_op": True,
                },
                artifact_refs=(artifact_id,),
                diff_refs=(diff_id,),
            )
        backup_path = path.with_suffix(path.suffix + f".{before_version[:8]}.bak")
        temp_dir = Path(tempfile.mkdtemp(prefix="jarvis_excel_", dir=str(path.parent)))
        temp_path = temp_dir / path.name
        try:
            shutil.copy2(path, temp_path)
            workbook = load_workbook(temp_path, read_only=False, data_only=False)
            sheet = workbook[change["sheet"]]
            cell = sheet[change["cell"]]
            cell.value = STATUS_LABELS[desired_status]
            cell.fill = PatternFill(fill_type="solid", fgColor=STATUS_COLORS[desired_status])
            workbook.save(temp_path)
            workbook.close()
            verify = load_workbook(temp_path, read_only=False, data_only=False)
            verify_cell = verify[change["sheet"]][change["cell"]]
            verified_value = str(verify_cell.value or "")
            verified_color = verify_cell.fill.fgColor.rgb[-6:]
            verify.close()
            if parse_status(verified_value) != desired_status or verified_color != STATUS_COLORS[desired_status]:
                return _failure(
                    ErrorCode.UNKNOWN,
                    "Saved workbook verification failed",
                    "excel.apply",
                )
            if not backup_path.exists():
                shutil.copy2(path, backup_path)
            os.replace(temp_path, path)
            after_version = self.fingerprint(path)
            artifact_id = f"artifact_{uuid4().hex}"
            diff_id = f"diff_{uuid4().hex}"
            artifact = Artifact(
                id=artifact_id,
                artifact_type="excel_workbook",
                display_name=path.name,
                uri_or_local_ref=str(path),
                mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                preview_ref=f"preview:{artifact_id}",
                version=after_version,
            )
            diff = ArtifactDiff(
                id=diff_id,
                artifact_id=artifact_id,
                before_version=before_version,
                after_version=after_version,
                changes=(change,),
                visual_highlights=(
                    {
                        "sheet": change["sheet"],
                        "cell": change["cell"],
                        "color": change["after_color"],
                    },
                ),
            )
            preview = self.preview(path, change["sheet"], highlight_cells=[change["cell"]])
            return ToolResult(
                status=ResultStatus.SUCCESS,
                summary=(
                    "Roadmap already had the requested status; workbook verified"
                    if proposal.get("already_applied")
                    else f"Updated {change['sheet']}!{change['cell']} and verified the workbook"
                ),
                data={
                    "artifact": artifact.model_dump(mode="json"),
                    "diff": diff.model_dump(mode="json"),
                    "preview": preview,
                    "workbook_path": str(path),
                },
                artifact_refs=(artifact_id,),
                diff_refs=(diff_id,),
                rollback_ref=str(backup_path),
            )
        except PermissionError as exc:
            return _failure(ErrorCode.PATH_NOT_ALLOWED, str(exc), "excel.apply")
        except Exception as exc:  # noqa: BLE001 - connector boundary
            return _failure(ErrorCode.UNKNOWN, f"{type(exc).__name__}: {exc}", "excel.apply")
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def preview(
        self,
        path: Path,
        sheet_name: str,
        highlight_cells: list[str] | None = None,
        max_rows: int = 30,
        max_cols: int = 12,
    ) -> dict[str, Any]:
        workbook = load_workbook(path, read_only=False, data_only=False)
        sheet = workbook[sheet_name]
        rows: list[list[dict[str, Any]]] = []
        highlight_set = set(highlight_cells or [])
        for row in sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row, max_rows),
            min_col=1,
            max_col=min(sheet.max_column, max_cols),
        ):
            rendered: list[dict[str, Any]] = []
            for cell in row:
                color = cell.fill.fgColor.rgb[-6:] if cell.fill and cell.fill.fgColor.rgb else ""
                rendered.append(
                    {
                        "coordinate": cell.coordinate,
                        "value": cell.value,
                        "fill": color,
                        "highlight": cell.coordinate in highlight_set,
                    }
                )
            rows.append(rendered)
        workbook.close()
        return {
            "sheet": sheet_name,
            "rows": rows,
            "truncated": sheet.max_row > max_rows or sheet.max_column > max_cols,
        }

    def rollback(self, backup_path: str, target_path: str) -> ToolResult:
        try:
            backup = self.resolve_allowed(backup_path)
            target = self.resolve_allowed(target_path)
            if not backup.exists():
                raise FileNotFoundError(backup)
            shutil.copy2(backup, target)
            return ToolResult(
                status=ResultStatus.SUCCESS,
                summary=f"Restored {target.name} from backup",
                data={"target_path": str(target), "fingerprint": self.fingerprint(target)},
            )
        except Exception as exc:  # noqa: BLE001
            return _failure(ErrorCode.UNKNOWN, f"Rollback failed: {exc}", "excel.rollback")


def parse_status(value: str) -> RoadmapStatus | None:
    normalized = _normalize(value).replace("-", " ").replace("_", " ")
    return STATUS_ALIASES.get(normalized)


def action_similarity(query: str, candidate: str) -> float:
    left = _normalize(query)
    right = _normalize(candidate)
    sequence = SequenceMatcher(None, left, right).ratio()
    left_tokens = set(left.split())
    right_tokens = set(right.split())
    token_score = len(left_tokens & right_tokens) / max(len(left_tokens | right_tokens), 1)
    containment = 1.0 if left in right or right in left else 0.0
    return sequence * 0.55 + token_score * 0.35 + containment * 0.10


def _normalize(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value.lower().strip())
    ascii_value = "".join(char for char in decomposed if not unicodedata.combining(char))
    cleaned = "".join(char if char.isalnum() else " " for char in ascii_value)
    return " ".join(cleaned.split())


def _find_header(headers: dict[str, int], aliases: set[str]) -> int | None:
    normalized_aliases = {_normalize(alias) for alias in aliases}
    for header, column in headers.items():
        if header in normalized_aliases:
            return column
    return None


def _failure(
    code: ErrorCode,
    message: str,
    component: str,
    retryable: bool = False,
) -> ToolResult:
    return ToolResult(
        status=ResultStatus.FAILURE,
        summary=message,
        diagnostics=(
            Diagnostic(
                code=code,
                message=message,
                component=component,
                retryable=retryable,
            ),
        ),
    )
