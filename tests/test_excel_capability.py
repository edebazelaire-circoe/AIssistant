from __future__ import annotations

import asyncio
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from jarvis_agent.application.capabilities import (
    CapabilityEngine,
    CapabilityRegistry,
    UpdateRoadmapWorkflow,
    register_excel_tools,
    register_update_roadmap_capability,
)
from jarvis_agent.application.event_store import EventStore
from jarvis_agent.application.tools import ToolRegistry, ToolRunner
from jarvis_agent.connectors.excel_roadmap import ExcelRoadmapConnector
from jarvis_agent.domain.models import CapabilityPolicy, PolicyMode, ResultStatus, RoadmapStatus


def create_workbook(path: Path, duplicate: bool = False) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Roadmap"
    sheet.append(["ID", "Action", "Responsable", "Mois", "Statut", "Formule"])
    sheet.append(["1", "Sécuriser le portail client", "Alice", "Juillet", "En cours", "=1+1"])
    if duplicate:
        sheet.append(["2", "Sécuriser le portail client", "Bob", "Juillet", "À faire", "=2+2"])
    sheet["E2"].fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    workbook.save(path)
    return path


def build_engine(tmp_path: Path):
    store = EventStore(tmp_path / "events.sqlite3")
    connector = ExcelRoadmapConnector(tmp_path)
    tools = ToolRegistry()
    register_excel_tools(tools, connector)
    runner = ToolRunner(tools, store)
    capabilities = CapabilityRegistry()
    workflow = UpdateRoadmapWorkflow(connector, runner)
    register_update_roadmap_capability(capabilities, workflow)
    engine = CapabilityEngine(capabilities, store)
    return store, connector, engine


def test_excel_safe_update_preserves_unrelated_formula(tmp_path: Path) -> None:
    path = create_workbook(tmp_path / "roadmap.xlsx")
    connector = ExcelRoadmapConnector(tmp_path)
    proposal = connector.propose(path, "Sécuriser le portail client", RoadmapStatus.DONE)
    result = connector.apply(proposal)
    assert result.status == ResultStatus.SUCCESS
    workbook = load_workbook(path, data_only=False)
    assert workbook["Roadmap"]["E2"].value == "Terminé"
    assert workbook["Roadmap"]["E2"].fill.fgColor.rgb.endswith("D9EAD3")
    assert workbook["Roadmap"]["F2"].value == "=1+1"
    workbook.close()
    assert result.data["diff"]["changes"][0]["cell"] == "E2"
    assert Path(result.rollback_ref).exists()


def test_excel_rejects_path_traversal(tmp_path: Path) -> None:
    connector = ExcelRoadmapConnector(tmp_path / "allowed")
    try:
        connector.resolve_allowed(tmp_path / "outside.xlsx")
    except PermissionError:
        pass
    else:
        raise AssertionError("outside path should be rejected")


def test_excel_ambiguous_row_is_not_mutated(tmp_path: Path) -> None:
    path = create_workbook(tmp_path / "ambiguous.xlsx", duplicate=True)
    connector = ExcelRoadmapConnector(tmp_path)
    try:
        connector.propose(path, "Sécuriser le portail client", RoadmapStatus.DONE)
    except RuntimeError as exc:
        assert "Ambiguous" in str(exc)
    else:
        raise AssertionError("duplicate rows must be ambiguous")
    workbook = load_workbook(path)
    assert workbook["Roadmap"]["E2"].value == "En cours"
    workbook.close()


def test_manual_policy_requires_single_use_approval(tmp_path: Path) -> None:
    create_workbook(tmp_path / "roadmap.xlsx")
    store, _, engine = build_engine(tmp_path)

    async def scenario() -> None:
        result = await engine.request(
            session_id="s",
            meeting_session_id="m",
            capability_id="update_roadmap",
            payload={"action_text": "Sécuriser le portail client", "desired_status": "done"},
            reason="test",
        )
        assert result.status == ResultStatus.NEEDS_APPROVAL
        approved = await engine.approve(result.data["approval_id"], result.data["approval_token"])
        assert approved.status == ResultStatus.SUCCESS
        reused = await engine.approve(result.data["approval_id"], result.data["approval_token"])
        assert reused.status == ResultStatus.FAILURE

    asyncio.run(scenario())
    event_types = [event.event_type for event in store.list_events("s")]
    assert "approval.requested" in event_types
    assert "approval.granted" in event_types
    assert "artifact.created" in event_types


def test_automatic_and_disabled_policies(tmp_path: Path) -> None:
    create_workbook(tmp_path / "roadmap.xlsx")
    _, _, engine = build_engine(tmp_path)

    async def scenario() -> None:
        await engine.set_policy(
            CapabilityPolicy(capability_id="update_roadmap", mode=PolicyMode.AUTOMATIC), "s"
        )
        automatic = await engine.request(
            session_id="s",
            meeting_session_id="m",
            capability_id="update_roadmap",
            payload={"action_text": "Sécuriser le portail client", "desired_status": "done"},
            reason="auto",
        )
        assert automatic.status == ResultStatus.SUCCESS
        await engine.set_policy(
            CapabilityPolicy(capability_id="update_roadmap", mode=PolicyMode.DISABLED), "s"
        )
        disabled = await engine.request(
            session_id="s",
            meeting_session_id="m",
            capability_id="update_roadmap",
            payload={"action_text": "Sécuriser le portail client", "desired_status": "in_progress"},
            reason="disabled",
        )
        assert disabled.status == ResultStatus.FAILURE
        assert disabled.diagnostics[0].code.value == "capability_disabled"

    asyncio.run(scenario())


def test_excel_repeated_update_is_noop_and_keeps_fingerprint(tmp_path: Path) -> None:
    path = create_workbook(tmp_path / "roadmap.xlsx")
    connector = ExcelRoadmapConnector(tmp_path)
    first = connector.apply(
        connector.propose(path, "Sécuriser le portail client", RoadmapStatus.DONE)
    )
    assert first.status == ResultStatus.SUCCESS
    fingerprint = connector.fingerprint(path)
    second_proposal = connector.propose(path, "Sécuriser le portail client", RoadmapStatus.DONE)
    assert second_proposal["already_applied"] is True
    second = connector.apply(second_proposal)
    assert second.status == ResultStatus.SUCCESS
    assert second.data["no_op"] is True
    assert connector.fingerprint(path) == fingerprint
