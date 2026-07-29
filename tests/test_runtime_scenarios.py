from __future__ import annotations

import asyncio
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from jarvis_agent.application.config import AppConfig
from jarvis_agent.application.runtime import JarvisRuntime, parse_roadmap_update
from jarvis_agent.domain.models import AgentState, PolicyMode, ResultStatus


def create_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Roadmap"
    sheet.append(["Action", "Statut"])
    sheet.append(["Sécuriser le portail client", "En cours"])
    sheet["B2"].fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    workbook.save(path)


def test_parse_roadmap_update_french() -> None:
    parsed = parse_roadmap_update("L'action Sécuriser le portail client est terminée.")
    assert parsed is not None
    assert parsed["action_text"] == "Sécuriser le portail client"
    assert parsed["desired_status"] == "done"


def test_reference_scenario_manual(tmp_path: Path) -> None:
    create_workbook(tmp_path / "workspace" / "roadmap.xlsx")
    runtime = JarvisRuntime(
        AppConfig(tmp_path / "data", tmp_path / "workspace", "127.0.0.1", 8765)
    )

    async def scenario() -> None:
        await runtime.initialize()
        assert (await runtime.set_state(AgentState.STANDBY)).status == ResultStatus.SUCCESS
        await runtime.start_meeting("Test")
        await runtime.inject_transcript(
            "Décision : lancer le pilote en septembre.", "Alice", auto_route=False
        )
        await runtime.inject_transcript(
            "Bob doit préparer le guide opérateur.", "Bob", auto_route=False
        )
        result = await runtime.inject_transcript(
            "L'action Sécuriser le portail client est terminée.", "Alice", auto_route=True
        )
        routed = result.data["routed_result"]
        assert routed["status"] == "needs_approval"
        approved = await runtime.capability_engine.approve(
            routed["data"]["approval_id"], routed["data"]["approval_token"]
        )
        assert approved.status == ResultStatus.SUCCESS
        ended = await runtime.end_meeting()
        assert ended.data["summary"]["decisions"]
        assert ended.data["summary"]["actions"]

    asyncio.run(scenario())
    workbook = load_workbook(tmp_path / "workspace" / "roadmap.xlsx")
    assert workbook["Roadmap"]["B2"].value == "Terminé"
    workbook.close()
    event_types = [event.event_type for event in runtime.store.list_events(runtime.session_id)]
    assert "transcript.partial" in event_types
    assert "transcript.final" in event_types
    assert "meeting.fact_extracted" in event_types
    assert "artifact.created" in event_types
    assert "meeting.summary_created" in event_types


def test_reference_scenario_automatic(tmp_path: Path) -> None:
    create_workbook(tmp_path / "workspace" / "roadmap.xlsx")
    runtime = JarvisRuntime(
        AppConfig(tmp_path / "data", tmp_path / "workspace", "127.0.0.1", 8765)
    )

    async def scenario() -> None:
        await runtime.initialize()
        await runtime.set_state(AgentState.STANDBY)
        await runtime.start_meeting("Auto")
        await runtime.set_capability_policy("update_roadmap", PolicyMode.AUTOMATIC)
        result = await runtime.inject_transcript(
            "L'action Sécuriser le portail client est terminée.", "Alice", auto_route=True
        )
        assert result.data["routed_result"]["status"] == "success"

    asyncio.run(scenario())
    events = runtime.store.list_events(runtime.session_id)
    assert not any(event.event_type == "approval.requested" for event in events)
    assert any(event.event_type == "artifact.created" for event in events)


def test_mic_off_rejects_audio(tmp_path: Path) -> None:
    import base64
    import numpy as np
    from jarvis_agent.domain.models import AudioFrame

    runtime = JarvisRuntime(
        AppConfig(tmp_path / "data", tmp_path / "workspace", "127.0.0.1", 8765)
    )
    raw = np.zeros(1600, dtype="<i2").tobytes()
    frame = AudioFrame(
        sequence=1,
        timestamp_ms=1,
        pcm16_b64=base64.b64encode(raw).decode(),
        level=0,
    )
    result = asyncio.run(runtime.ingest_audio(frame))
    assert result.status == ResultStatus.CANCELLED
    assert runtime.audio_buffer.stats.frame_count == 0


def test_runtime_rebuilds_latest_meeting_projection(tmp_path: Path) -> None:
    runtime = JarvisRuntime(
        AppConfig(tmp_path / "data", tmp_path / "workspace", "127.0.0.1", 8765)
    )

    async def scenario() -> None:
        await runtime.initialize()
        await runtime.set_state(AgentState.STANDBY)
        await runtime.start_meeting("Recovery")
        await runtime.inject_transcript("Décision : garder le stockage local.", "Alice", auto_route=False)

    asyncio.run(scenario())
    restarted = JarvisRuntime(
        AppConfig(tmp_path / "data", tmp_path / "workspace", "127.0.0.1", 8765)
    )
    assert restarted.meeting is not None
    assert restarted.meeting.title == "Recovery"
    assert restarted.transcript[-1].text == "Décision : garder le stockage local."
    assert restarted.facts[-1].fact_type.value == "decision"
